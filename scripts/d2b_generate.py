"""
D2b — candidate learner responses written by DeepSeek, for the researcher to screen, edit and
label (evaluation_plan §4.2 as amended in V1.2: "LLM-generated, researcher-screened and annotated").

Run from the project root (needs DEEPSEEK_API_KEY):
    python scripts/d2b_generate.py                 # 60 candidates into the empty rows of eval/d2b_cases.xlsx
    python scripts/d2b_generate.py --kind algebra --count 3   # top up one kind (see KINDS)

How independence is kept
    * DeepSeek is asked to play a 初一 student. It sees the problem text and a plain teacher's
      description of the kind of response wanted ("solved correctly but slipped in one step"). It
      never sees the diagnosis rules, the error codes, the reference equation or the answer.
    * Only the RESPONSE columns are filled. Every label column stays empty for the researcher.
    * The intended kind of each candidate is NOT written into the workbook (it would anchor the
      labels). It is kept in eval/d2b_generation_log.json, together with the exact text generated,
      so d2b_import.py can record whether each final case was kept as generated, edited, or
      written by the researcher.
    * Rows are shuffled, so the order does not give the intended kind away.
    * Nothing is filtered with the system's own checks: a candidate that turns out to be something
      other than intended is simply labelled as what it is (or deleted).
"""
import argparse
import json
import random
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import d2b_io as io  # noqa: E402

LOG = io.ROOT / "eval" / "d2b_generation_log.json"
SEED = 20260930
PROMPT_VERSION = "d2b-gen-v1"

SYSTEM = """你在帮助一项数学教育研究收集“初中生作答样本”。请扮演一名初一学生，针对给出的应用题写出作答。
要求：
- 写得像真实学生写在作业本上的样子：可以口语化，可以用全角符号（＋ － × ÷ ＝）、省略乘号（2x）、写“解：设…”。
- 按“作答类型”的描述来写。如果描述要求出错，只在描述的地方出错，其他部分照常写；出错后按错的结果继续往下算，不要自己改正，也不要注明哪里错了。
- 不要写任何解释、批注或老师的话，只写学生写下的内容。
- 只输出 JSON，字段见“需要写的内容”。"""

WHAT_TO_WRITE = {
    "modelling": ('{"variable_definition": "学生写的设未知数那句", "equation": "学生列的方程（一个）"}',
                  "只写设未知数和列出的方程，不用解。"),
    "execution": ('{"variable_definition": "设…", "equation": "列出的方程", "steps": ["解方程的第1步", "第2步", "…", "x = …"]}',
                  "写设未知数、方程，然后一步一行地解方程（steps 从方程变形的第一步开始，不要重抄方程），解到 x = … 为止。"),
    "answer": ('{"final_answer_text": "学生写的最终答案，一句话"}',
               "只写最后的答案那一句（例如“答：……”），不写过程。"),
}

# kind: (layer, teacher's description, question ids that suit it or None for any)
KINDS = {
    "goal_answer": ("answer", "没看清题目问的是什么：回答成了题目里别的量，或者题目问两个量只回答了其中一个。", None),
    "goal_model": ("modelling", "没看清题目要求的量：设未知数和列方程时求的是题目里另一个量，而不是题目问的那个。", None),
    "define": ("modelling", "设未知数时说得不清楚或者前后不一致：比如只写“设x”不说是什么，或者设的是一个量、方程里却把 x 当成另一个量来用。", None),
    "relation": ("modelling", "数量关系列错了：比如该加的写成了减、漏掉了一部分数量、把总数和其中一部分弄混。", None),
    "language": ("modelling", "把题目里的中文说法翻译成算式时理解错了，比如“比…多/少/贵”的方向弄反、“打八折”、“涨价20%”、“是…的几分之几/几倍”这类说法理解错。",
                 ["Q01", "Q02", "Q03", "Q06", "Q07", "Q08", "Q11", "Q12", "Q14"]),
    "unit_answer": ("answer", "单位出了问题：最终答案的单位写错，或者没有做单位换算。", ["Q04", "Q09", "Q13", "Q10"]),
    "unit_model": ("modelling", "单位出了问题：列方程时把不同单位的量直接放在一起算（比如分钟和小时、米和千米混用）。", ["Q04", "Q09", "Q13"]),
    "arithmetic": ("execution", "方程列对了，但解的过程中有一步算术算错了（加减乘除的计算错误），后面照着错的结果继续算。", None),
    "algebra": ("execution", "方程列对了，但解的过程中有一步变形出错，比如移项忘了变号、去括号时只乘了括号里的第一项、只对等式一边做了运算，后面照着错的结果继续算。", None),
    "multi": ("execution", "列方程时有一处错误，而且解方程时又有一步算错或变形出错（两处独立的错误）。", None),
    "correct_equiv": ("modelling", "完全正确，但方程的写法和课本常见写法不同（比如项的顺序不同、移到另一边、写成等价的另一种形式）。", None),
    "correct_altdef": ("modelling", "完全正确，但设的未知数和常见做法不同（比如题目有两个量时设另一个量为 x）。",
                       ["Q01", "Q02", "Q05", "Q06", "Q11", "Q12", "Q13", "Q15"]),
    "correct_altwork": ("execution", "完全正确，但解方程的步骤和常见做法不一样（比如先两边同除再移项、几步合成一步、先去分母）。", None),
    "correct_answer": ("answer", "只写了最终答案，答案完全正确，单位也对。", None),
    "ambiguous": ("modelling", "方程里用了有歧义、让老师看不清的写法，比如把 x/3 写成 1/3x、把分数写成 x/4+5 却可能是 x/(4+5) 的意思。", None),
    "answer_unclear": ("answer", "只写了最终答案，答案不对，但光看这一个答案看不出是哪里出了错。", None),
}

# 60 slots — meets evaluation_plan §4.2 if the labels come out roughly as intended; the
# composition check in d2b_import.py uses the researcher's labels, not these intentions.
PLAN = ([("goal_answer", 3), ("goal_model", 3), ("define", 5), ("relation", 5), ("language", 6),
         ("unit_answer", 2), ("unit_model", 4), ("arithmetic", 5), ("algebra", 5), ("multi", 5),
         ("correct_equiv", 4), ("correct_altdef", 3), ("correct_altwork", 3), ("correct_answer", 2),
         ("ambiguous", 3), ("answer_unclear", 2)])


def slots(seed=SEED):
    rng = random.Random(seed)
    bank_ids = list(io.load_bank())
    out = []
    for kind, n in PLAN:
        pool = KINDS[kind][2] or bank_ids
        picks = rng.sample(pool, n) if n <= len(pool) else [rng.choice(pool) for _ in range(n)]
        out += [(kind, q) for q in picks]
    return out


def prompt(question: dict, kind: str) -> str:
    layer, desc, _ = KINDS[kind]
    fmt, what = WHAT_TO_WRITE[layer]
    return (f"应用题：{question['question']}\n\n作答类型：{desc}\n\n需要写的内容：{what}\n"
            f"输出格式：{fmt}")


def parse(raw: str, layer: str) -> dict | None:
    try:
        d = json.loads(raw)
    except Exception:  # noqa: BLE001
        return None
    if not isinstance(d, dict):
        return None
    out = {}
    if layer in ("modelling", "execution"):
        out["variable_definition"] = str(d.get("variable_definition") or "").strip()
        out["equation"] = str(d.get("equation") or "").strip()
        if not out["equation"]:
            return None
    if layer == "execution":
        steps = d.get("steps") or []
        if isinstance(steps, str):
            steps = io.split_steps(steps)
        out["steps"] = "\n".join(str(s).strip() for s in steps if str(s).strip())
        if not out["steps"]:
            return None
    if layer == "answer":
        out["final_answer_text"] = str(d.get("final_answer_text") or "").strip()
        if not out["final_answer_text"]:
            return None
    return out


def empty_rows(ws):
    header = [c.value for c in ws[1]]
    qcol = header.index("question_id") + 1
    return header, [r for r in range(2, ws.max_row + 1) if not ws.cell(row=r, column=qcol).value]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=str(io.WORKBOOK))
    ap.add_argument("--log", default=str(LOG))
    ap.add_argument("--kind", choices=sorted(KINDS), help="generate only this kind (a top-up)")
    ap.add_argument("--count", type=int, default=3, help="with --kind: how many")
    ap.add_argument("--seed", type=int, default=SEED)
    ap.add_argument("--temperature", type=float, default=1.0)
    args = ap.parse_args()

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment

    from core.llm_client import DeepSeekBackend, LLMError
    try:
        backend = DeepSeekBackend(temperature=args.temperature)
    except LLMError as exc:
        sys.exit(f"✗ {exc}")

    bank = io.load_bank()
    if args.kind:
        pool = KINDS[args.kind][2] or list(bank)
        rng = random.Random(f"{args.seed}:{args.kind}:{datetime.now():%Y%m%d%H%M%S}")
        plan = [(args.kind, rng.choice(pool)) for _ in range(args.count)]
    else:
        plan = slots(args.seed)
    wb_path = Path(args.workbook)
    wb = load_workbook(wb_path)
    ws = wb[io.CASE_SHEET]
    header, free = empty_rows(ws)
    if len(free) < len(plan):
        sys.exit(f"✗ only {len(free)} empty row(s) in {wb_path.name}, {len(plan)} needed")
    log_path = Path(args.log)
    log = json.loads(log_path.read_text(encoding="utf-8")) if log_path.exists() else {"runs": [], "cases": {}}

    generated = []
    for i, (kind, qid) in enumerate(plan, 1):
        layer = KINDS[kind][0]
        resp = None
        for _ in range(3):
            try:
                resp = parse(backend.complete(SYSTEM, prompt(bank[qid], kind)), layer)
            except LLMError as exc:
                print(f"  ! {qid} {layer}: {exc}")
            if resp:
                break
        print(f"  {i:>2}/{len(plan)} {qid} {layer:<10} {'ok' if resp else 'FAILED'}")
        if resp:
            generated.append((kind, qid, layer, resp))

    random.Random(f"{args.seed}:rows").shuffle(generated)
    col = {h: j + 1 for j, h in enumerate(header)}
    run_id = datetime.now().strftime("%Y%m%d-%H%M%S")
    for row, (kind, qid, layer, resp) in zip(free, generated):
        cid = ws.cell(row=row, column=col["case_id"]).value or f"D2b-{row - 1:03d}"
        ws.cell(row=row, column=col["case_id"], value=cid)
        ws.cell(row=row, column=col["question_id"], value=qid)
        ws.cell(row=row, column=col["layer"], value=layer)
        for k, v in resp.items():
            c = ws.cell(row=row, column=col[k], value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=col["review_status"], value="pending")
        log["cases"][cid] = {"run": run_id, "intended_kind": kind, "question_id": qid, "layer": layer,
                             "generated": resp}
    log["runs"].append({"run": run_id, "prompt_version": PROMPT_VERSION, "backend": backend.describe(),
                        "seed": args.seed, "n_slots": len(plan), "n_generated": len(generated)})
    wb.save(wb_path)
    log_path.write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n✓ {len(generated)} candidate(s) written into {wb_path.name}; intentions kept in {log_path.name}"
          " (do not open it before you have labelled)")


if __name__ == "__main__":
    main()
