"""
D2b — shared helpers: the annotation workbook (eval/d2b_cases.xlsx) and its conversion to
diagnostic-case records (evaluation_plan §4.2–4.3, error_taxonomy §14.2).

Used by d2b_template.py, d2b_import.py, d2b_second.py and eval_diagnosis_d2b.py.
Nothing here looks at the system's diagnosis: the reference labels come only from the annotator.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.normalizer import normalize_line  # noqa: E402

WORKBOOK = ROOT / "eval" / "d2b_cases.xlsx"
CASES_JSON = ROOT / "eval" / "diagnostic_cases.json"
BANK_JSON = ROOT / "data" / "question_bank.json"
SECOND_WORKBOOK = ROOT / "eval" / "d2b_second_annotator.xlsx"
SECOND_JSON = ROOT / "eval" / "d2b_second_annotation.json"
DISAGREEMENTS_MD = ROOT / "eval" / "annotation_disagreements.md"
CORRECTIONS_MD = ROOT / "eval" / "annotation_corrections.md"

CASE_SHEET = "cases"

LAYERS = ["modelling", "execution", "answer"]
STATE_OF_LAYER = {"modelling": "S2", "execution": "S3", "answer": "S4"}
STATUSES = ["correct", "incorrect", "uncertain", "not_assessable"]
CATS = [f"E{i}" for i in range(1, 8)]
SUBTYPES = [f"E4-0{i}" for i in range(1, 9)]
LEVELS = ["A", "B", "C", "D"]
REVIEW = ["pending", "reviewed", "adjudicated"]
TAGS = ["equivalent_equation", "alt_definition", "alt_working", "ambiguous_notation"]

# (key, width, what the annotator writes) — the header row holds the keys exactly
COLUMNS = [
    ("case_id", 11, "编号，已预填 D2b-001…；不用的行留着即可（没有 question_id 的行会被跳过）"),
    ("question_id", 11, "对应题库哪一题（Q01–Q15，下拉选择；题目见“题目”工作表）"),
    ("layer", 12, "modelling = 列方程（S2）；execution = 解方程过程（S3）；answer = 只看最终答案（S4）"),
    ("variable_definition", 26, "学生写的“设…”（modelling 可填；execution 可填；不写就留空）"),
    ("equation", 26, "modelling：学生列的方程（必填）。execution：学生在 S2 列好的方程（选填；填了就从它开始检查第 1 行）"),
    ("steps", 30, "execution：学生的解方程过程，一行一步（单元格内换行 Alt+Enter，或用 || 分隔）"),
    ("final_answer_text", 22, "answer：学生写的最终答案原文，例如“小明15岁，小华10岁”"),
    ("reference_status", 15, "你的判断：correct / incorrect / uncertain（两种解释都说得通）/ not_assessable（看不懂）"),
    ("primary_error", 12, "incorrect 时的主要错误类别 E1–E7（下拉）"),
    ("error_subtype", 12, "只在 E4 时填 E4-01…E4-08（选填）"),
    ("secondary_errors", 14, "同时存在的其他独立错误，例如 E6;E7（多错误样本用）"),
    ("first_invalid_step", 12, "execution 且有错时：steps 里第几行是第一个错的（从 1 数，只数 steps 单元格里的行）"),
    ("evidence_level", 10, "A 直接可见 / B 只从答案推断 / C 部分证据 / D 证据不足"),
    ("tags", 22, "选填，分号分隔：equivalent_equation; alt_definition; alt_working; ambiguous_notation"),
    ("evidence", 30, "作答里哪一处支持你的标注（原文摘录即可）"),
    ("note", 24, "备注（选填）"),
    ("review_status", 12, "pending（写完还没复查）→ reviewed（复查定稿）；全部 reviewed 后才能冻结和正式运行"),
]
KEYS = [c[0] for c in COLUMNS]
REFERENCE_FIELDS = ("reference_status", "primary_error", "error_subtype", "secondary_errors", "first_invalid_step")


def load_bank(path: Path = BANK_JSON) -> dict:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return {q["question_id"]: q for q in data["questions"]}


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_workbook(path: Path, sheet: str = CASE_SHEET) -> list[dict]:
    """Rows as {key: str}, with the Excel row number in '_row'. Rows without question_id are skipped."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"✗ {path.name}: no sheet named '{sheet}'")
    ws = wb[sheet]
    header = [_cell(c.value) for c in ws[1]]
    missing = [k for k in ("case_id", "question_id") if k not in header]
    if missing:
        raise SystemExit(f"✗ {path.name}: header row lacks {missing}")
    rows = []
    for i, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {h: _cell(v) for h, v in zip(header, values) if h}
        if not row.get("question_id"):
            continue
        row["_row"] = i
        rows.append(row)
    return rows


def split_steps(text: str) -> list[str]:
    parts = re.split(r"\r?\n|\|\|", text or "")
    return [p.strip() for p in parts if p.strip()]


def parse_codes(text: str) -> list[str]:
    return sorted(set(re.findall(r"E[1-7]", (text or "").upper())))


def parse_tags(text: str) -> list[str]:
    return sorted({t.strip() for t in re.split(r"[;；,，\s]+", text or "") if t.strip()})


def _norm(line):
    n = normalize_line(line)
    return n.normalized if n.ok else None


def validate(rows: list[dict], bank: dict) -> tuple[list[str], list[str]]:
    """(errors, warnings). Errors stop the import; warnings are for the annotator to look at."""
    errors, warnings = [], []
    seen = {}
    for r in rows:
        where = f"row {r['_row']} ({r.get('case_id') or 'no case_id'})"
        cid = r.get("case_id")
        if not cid:
            errors.append(f"{where}: case_id is empty")
        elif cid in seen:
            errors.append(f"{where}: case_id repeats row {seen[cid]}")
        else:
            seen[cid] = r["_row"]
        if r["question_id"] not in bank:
            errors.append(f"{where}: question_id '{r['question_id']}' is not in the question bank")
        layer, status = r.get("layer", ""), r.get("reference_status", "")
        if layer not in LAYERS:
            errors.append(f"{where}: layer must be one of {LAYERS}")
        if status not in STATUSES:
            errors.append(f"{where}: reference_status must be one of {STATUSES}")
        steps = split_steps(r.get("steps", ""))
        if layer == "modelling" and not r.get("equation"):
            errors.append(f"{where}: a modelling case needs an equation")
        if layer == "execution" and not steps:
            errors.append(f"{where}: an execution case needs steps")
        if layer == "answer" and not r.get("final_answer_text"):
            errors.append(f"{where}: an answer case needs final_answer_text")
        if layer != "execution" and steps:
            warnings.append(f"{where}: steps are ignored for a {layer} case")
        if layer == "answer" and r.get("equation"):
            warnings.append(f"{where}: equation is ignored for an answer case")

        prim, sub = r.get("primary_error", "").upper(), r.get("error_subtype", "").upper()
        if prim and prim not in CATS:
            errors.append(f"{where}: primary_error must be E1–E7")
        if status == "incorrect" and not prim:
            warnings.append(f"{where}: incorrect without primary_error (allowed only when the evidence permits no category)")
        if status in ("correct", "not_assessable") and prim:
            errors.append(f"{where}: a {status} case cannot have a primary_error")
        if sub and sub not in SUBTYPES:
            errors.append(f"{where}: error_subtype must be E4-01…E4-08")
        if sub and prim != "E4":
            errors.append(f"{where}: error_subtype is only for primary_error E4")
        sec = parse_codes(r.get("secondary_errors", ""))
        if sec and status != "incorrect":
            errors.append(f"{where}: secondary_errors only on incorrect cases")
        if prim and prim in sec:
            warnings.append(f"{where}: primary_error is repeated in secondary_errors")

        fis = r.get("first_invalid_step", "")
        if fis:
            if not fis.isdigit():
                errors.append(f"{where}: first_invalid_step must be a whole number")
            elif layer != "execution":
                errors.append(f"{where}: first_invalid_step only for execution cases")
            elif not 1 <= int(fis) <= len(steps):
                errors.append(f"{where}: first_invalid_step {fis} but steps has {len(steps)} line(s)")
            elif status == "correct":
                errors.append(f"{where}: a correct case has no first_invalid_step")
        elif layer == "execution" and status == "incorrect" and prim in ("E6", "E7"):
            warnings.append(f"{where}: execution error without first_invalid_step (not counted for step localisation)")

        if r.get("evidence_level") and r["evidence_level"].upper() not in LEVELS:
            errors.append(f"{where}: evidence_level must be A/B/C/D")
        unknown = [t for t in parse_tags(r.get("tags", "")) if t not in TAGS]
        if unknown:
            warnings.append(f"{where}: unknown tag(s) {unknown} — kept, but not counted in the composition check")
        if r.get("review_status", "pending") not in REVIEW + [""]:
            errors.append(f"{where}: review_status must be one of {REVIEW}")
        unread = [i + 1 for i, s in enumerate(steps) if _norm(s) is None]
        if unread and "ambiguous_notation" not in parse_tags(r.get("tags", "")):
            warnings.append(f"{where}: steps line(s) {unread} cannot be parsed — fine if intended, then tag ambiguous_notation")
    return errors, warnings


def to_record(r: dict, split: str) -> dict:
    """One reference-annotated diagnostic case (error_taxonomy §14.2 + source/split)."""
    layer, status = r["layer"], r["reference_status"]
    prim = r.get("primary_error", "").upper() or None
    response = {}
    if r.get("variable_definition") and layer != "answer":
        response["variable_definition"] = r["variable_definition"]
    if layer == "modelling":
        response["equation"] = r["equation"]
    elif layer == "execution":
        if r.get("equation"):
            response["equation"] = r["equation"]
        response["steps"] = split_steps(r["steps"])
    else:
        response["final_answer_text"] = r["final_answer_text"]
    normalized = {k: (v if k in ("variable_definition", "final_answer_text")
                      else [_norm(s) for s in v] if isinstance(v, list) else _norm(v))
                  for k, v in response.items()}
    diagnosis_status = {"correct": "no_error_detected", "uncertain": "uncertain",
                        "not_assessable": "insufficient_evidence"}.get(status, "diagnosed" if prim else "uncertain")
    fis = r.get("first_invalid_step", "")
    return {
        "case_id": r["case_id"], "question_id": r["question_id"], "source": "hand_crafted", "split": split,
        "layer": layer, "state": STATE_OF_LAYER[layer],
        "learner_response": response, "normalized_response": normalized,
        "reference_status": status, "primary_error": prim,
        "error_subtype": r.get("error_subtype", "").upper() or None,
        "candidate_errors": [prim] if prim else [],
        "secondary_errors": parse_codes(r.get("secondary_errors", "")),
        "evidence_level": r.get("evidence_level", "").upper() or None,
        "diagnosis_status": diagnosis_status,
        "evidence": r.get("evidence") or None,
        "first_invalid_step": int(fis) if fis else None,
        "tags": parse_tags(r.get("tags", "")),
        "annotation_note": r.get("note") or None,
        "review_status": r.get("review_status") or "pending",
    }


def stratum(case: dict) -> str:
    return case["primary_error"] if case["reference_status"] == "incorrect" and case["primary_error"] \
        else case["reference_status"]


def question_rows(bank: dict) -> list[tuple]:
    return [(q["question_id"], q["topic"], q["question"], q["variable_definition"], q["equation"],
             q["final_answer_text"]) for q in bank.values()]
