"""
D2b — create the blank annotation workbook eval/d2b_cases.xlsx (evaluation_plan §4.2).

Run from the project root:
    python scripts/d2b_template.py            # refuses to overwrite an existing workbook
    python scripts/d2b_template.py --rows 60

Sheets: cases (fill in, with drop-down lists), 说明 (how to fill it in and what the set must
contain), 题目 (the 15 bank questions the cases refer to).
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import d2b_io as io  # noqa: E402

GUIDE = [
    ("怎么填", ""),
    ("1", "每一行是一个“学生作答”。先选 question_id 和 layer，再照学生会写的样子填作答内容，最后填你的判断（reference_status 及之后的列）。"),
    ("2", "作答要像真实学生写的：口语化的“设”、全角符号、2x＋5＝25、漏写乘号都可以。不要照着系统的变异模板写。"),
    ("3", "判断只根据作答本身（error_taxonomy.md §14.3）：看得见的证据，不猜意图；两种解释都说得通就标 uncertain；只有最终答案时不要硬推类别。"),
    ("4", "写的时候不要运行系统、不要看诊断代码（core/diagnosis.py）。标注必须先于系统运行完成（评估计划 §5.1）。"),
    ("5", "execution 行：steps 一行一步，单元格内换行用 Alt+Enter（也可以用 || 分隔）。first_invalid_step 从 1 数，只数 steps 单元格里的行。"),
    ("6", "execution 行的 equation 选填：填了表示学生在 S2 已经列好这个方程，steps 可以不重抄它，第 1 行按“从这个方程变形”来检查。"),
    ("7", "全部写完、自己复查一遍后，把 review_status 改成 reviewed。之后改标签需要写明理由（见 d2b_import.py --allow-relabel）。"),
    ("8", "用 DeepSeek 生成候选作答时（python scripts/d2b_generate.py）：逐条看，不像真实学生的删掉或改写，然后像自己写的一样独立标注。生成意图存在 eval/d2b_generation_log.json，标注完成并冻结之前不要打开它。"),
    ("", ""),
    ("数据集必须包含（评估计划 §4.2，导入时会自动检查）", ""),
    ("总数", "40–60 条"),
    ("E1–E7", "每个类别至少 5 条（按 primary_error 计）"),
    ("正确作答", "占 20–30%：包括等价方程、另一种设法、另一种解法、只写对最终答案 — 用 tags 标 equivalent_equation / alt_definition / alt_working"),
    ("uncertain", "至少 5 条（reference_status = uncertain）"),
    ("记号有歧义", "至少 3 条，例如 1/3x — tags 标 ambiguous_notation"),
    ("只有最终答案且不对", "至少 5 条（layer = answer，reference_status 不是 correct）"),
    ("多个错误", "至少 5 条（secondary_errors 不为空）"),
    ("逐行解方程", "至少 10 条（layer = execution，且填了 first_invalid_step）"),
    ("第二标注人", "至少 20% 由另一个人盲标（python scripts/d2b_second.py export）"),
    ("", ""),
    ("错误类别（详见 docs/error_taxonomy.md）", ""),
    ("E1", "没读懂题目要求什么（求错了量）"),
    ("E2", "未知数定义不清或前后不一致"),
    ("E3", "数量关系错误"),
    ("E4", "中文表述翻译成数学时的固定错误（E4-01 比…多 方向反 … E4-08 少了三分之一）"),
    ("E5", "单位或数量级错误"),
    ("E6", "算术计算错误"),
    ("E7", "代数变形或符号错误（移项不变号、去括号等）"),
    ("", ""),
    ("填完之后", ""),
    ("检查", "python scripts/d2b_import.py --check"),
    ("导入", "python scripts/d2b_import.py"),
    ("第二标注人", "python scripts/d2b_second.py export   然后   python scripts/d2b_second.py compare"),
    ("冻结", "所有 review_status = reviewed 后提交：git add eval  → git commit → git push"),
    ("运行", "python scripts/eval_diagnosis_d2b.py            （默认只跑 test，70%）"),
]


def build(path: Path, columns, bank: dict, prefill: list[dict], title_note: str, guide=GUIDE,
          validations=True):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = io.CASE_SHEET
    head_fill = PatternFill("solid", fgColor="DDE7F0")
    label_fill = PatternFill("solid", fgColor="FFF4D6")
    for j, (key, width, help_text) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=j, value=key)
        c.font = Font(bold=True)
        c.fill = label_fill if key in io.REFERENCE_FIELDS + ("evidence_level", "evidence", "review_status") else head_fill
        c.comment = Comment(help_text, "d2b")
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "C2"
    keys = [c[0] for c in columns]
    for i, row in enumerate(prefill, start=2):
        for key, value in row.items():
            if key in keys:
                ws.cell(row=i, column=keys.index(key) + 1, value=value)
    last = max(len(prefill) + 1, 2)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    if validations:
        def lst(key, options):
            if key in keys:
                col = ws.cell(row=1, column=keys.index(key) + 1).column_letter
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                dv.add(f"{col}2:{col}{last}")
                ws.add_data_validation(dv)
        lst("question_id", list(bank))
        lst("layer", io.LAYERS)
        lst("reference_status", io.STATUSES)
        lst("primary_error", io.CATS)
        lst("error_subtype", io.SUBTYPES)
        lst("evidence_level", io.LEVELS)
        lst("review_status", io.REVIEW)
        if "first_invalid_step" in keys:
            col = ws.cell(row=1, column=keys.index("first_invalid_step") + 1).column_letter
            dv = DataValidation(type="whole", operator="between", formula1="1", formula2="30", allow_blank=True)
            dv.add(f"{col}2:{col}{last}")
            ws.add_data_validation(dv)

    g = wb.create_sheet("说明")
    g.column_dimensions["A"].width = 30
    g.column_dimensions["B"].width = 110
    g.append([title_note, ""])
    g["A1"].font = Font(bold=True, size=13)
    g.append(["", ""])
    for a, b in guide:
        g.append([a, b])
        if a and not b:
            g.cell(row=g.max_row, column=1).font = Font(bold=True)
    g.append(["", ""])
    g.append(["列说明", ""])
    g.cell(row=g.max_row, column=1).font = Font(bold=True)
    for key, _, help_text in columns:
        g.append([key, help_text])
    for row in g.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    q = wb.create_sheet("题目")
    q.append(["question_id", "topic", "question", "标准设法", "标准方程", "标准答案"])
    for c in q[1]:
        c.font = Font(bold=True)
    for r in io.question_rows(bank):
        q.append(list(r))
    for col, w in zip("ABCDEF", (11, 16, 70, 30, 26, 30)):
        q.column_dimensions[col].width = w
    for row in q.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(io.WORKBOOK))
    ap.add_argument("--rows", type=int, default=60)
    ap.add_argument("--force", action="store_true", help="overwrite an existing workbook (loses what is in it)")
    args = ap.parse_args()
    out = Path(args.out)
    if out.exists() and not args.force:
        sys.exit(f"✗ {out} already exists — it may hold your annotations. Use --force only if you mean to replace it.")
    prefill = [{"case_id": f"D2b-{i:03d}", "review_status": "pending"} for i in range(1, args.rows + 1)]
    build(out, io.COLUMNS, io.load_bank(), prefill, "D2b 手写诊断样本 — 填写说明")
    print(f"✓ wrote {out} ({args.rows} blank rows)")


if __name__ == "__main__":
    main()
